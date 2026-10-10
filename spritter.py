import copy
from contextlib import closing
import glob
import json
import os
import re
import openpyxl
from openpyxl.styles import Protection
import pandas as pd

# ==========================================
# 1. 設定ファイル (JSON) の読み込み
# ==========================================
CONFIG_FILE = "config.json"

try:
  with open(CONFIG_FILE, "r", encoding="utf-8") as f:
    config = json.load(f)
except FileNotFoundError:
  raise FileNotFoundError(
      f"設定ファイル '{CONFIG_FILE}' が見つかりません。同じ階層に配置してください。"
  )

INPUT_FOLDER = config["input_folder"]
OUTPUT_FOLDER = config["output_folder"]
TARGET_SHEET_NAME = config["target_sheet_name"]
HEADER_ROW = config["header_row"]
TOOL_COL_IDX = config["tool_col_idx"]
TEXT_COL_INDICES = config["text_col_indices"]

CIRCLED_MAPPING = config["mapping"]
CIRCLED_NUMS = config["circled_nums"]
CIRCLED_TO_IDX = {num: i for i, num in enumerate(CIRCLED_NUMS)}
CIRCLED_CHARS = "".join(
    sorted({
        char
        for char in [*CIRCLED_MAPPING.keys(), *CIRCLED_NUMS]
        if len(char) == 1 and not char.isascii()
    })
)
ALL_CIRCLED_CHARS = f"[{re.escape(CIRCLED_CHARS)}]"
PREFIX_PATTERN = re.compile(
    rf"^((?:{ALL_CIRCLED_CHARS}\s*(?:[〜~-]\s*{ALL_CIRCLED_CHARS})?"
    rf"\s*[、,]?\s*)+)(?:[:：]\s*)?(.*)$"
)


# ==========================================
# 2. 丸数字標準化・解析関数
# ==========================================
def normalize_circled_number(char):
  return CIRCLED_MAPPING.get(char, char)


def expand_prefix_to_nums(prefix):
  """プレフィックス（例: '①〜③、⑤' や '①②③'）から

  該当する丸数字のリスト（重複なし・順序維持）を抽出・展開する関数
  """
  expanded = []
  parts = re.split(r"[、,]", prefix)

  for part in parts:
    part = part.strip()
    if not part:
      continue

    # 範囲指定のチェック（例: ①〜③）
    range_match = re.search(
        rf"({ALL_CIRCLED_CHARS})\s*[〜~-]\s*({ALL_CIRCLED_CHARS})", part
    )
    if range_match:
      start_char, end_char = range_match.groups()
      norm_start = normalize_circled_number(start_char)
      norm_end = normalize_circled_number(end_char)
      start_idx = CIRCLED_TO_IDX.get(norm_start, 0)
      end_idx = CIRCLED_TO_IDX.get(norm_end, 0)
      if start_idx <= end_idx:
        for idx in range(start_idx, end_idx + 1):
          if idx < len(CIRCLED_NUMS):
            num_char = CIRCLED_NUMS[idx]
            if num_char not in expanded:
              expanded.append(num_char)
      continue

    # 個別または連続する丸数字のチェック
    found_in_part = re.findall(rf"{ALL_CIRCLED_CHARS}", part)
    if found_in_part:
      for num_char in found_in_part:
        norm_char = normalize_circled_number(num_char)
        if norm_char in CIRCLED_TO_IDX:
          if norm_char not in expanded:
            expanded.append(norm_char)

  return expanded


def parse_cell_to_blocks(text):
  """セル内のテキストを解析し、

  {丸数字: [その丸数字が属するテキスト項目のリスト]} と、
  共通テキストを抽出する
  """
  if pd.isna(text) or not str(text).strip():
    return {}, []

  lines = [line.strip() for line in str(text).splitlines() if line.strip()]
  # 各丸数字に紐づくテキストブロックのリストを保持する辞書
  mapping = {num: [] for num in CIRCLED_NUMS}
  common_texts = []

  for line in lines:
    match = PREFIX_PATTERN.match(line)
    if not match:
      common_texts.append(line)
      continue

    prefix = match.group(1).strip()
    body = match.group(2).strip()
    if not body:
      body = line

    nums_in_line = expand_prefix_to_nums(prefix)

    if nums_in_line:
      for num in nums_in_line:
        mapping[num].append(body)
    else:
      common_texts.append(line)

  return mapping, common_texts


def format_cell_value(num, bodies, common_texts=None, include_num=True):
  """指定丸数字に紐づく本文と共通本文を1つのセル値にまとめる。"""
  values = []
  if bodies:
    values.extend(f"{num}：{body}" if include_num else f"{num}{body}" for body in bodies)
  if common_texts:
    values.extend(common_texts)
  return "\n".join(values) if values else "ー"


def fill_empty_cells(row):
  for col_idx in range(len(row)):
    val = row.iloc[col_idx]
    if pd.isna(val) or str(val).strip() == "":
      row.iloc[col_idx] = "ー"
  return row


def copy_cell_format(source_cell, target_cell):
  if source_cell.has_style:
    target_cell._style = copy.copy(source_cell._style)
  target_cell.number_format = source_cell.number_format
  target_cell.alignment = copy.copy(source_cell.alignment)
  target_cell.protection = copy.copy(source_cell.protection)


# ==========================================
# 3. 指定フォルダ内の複数ファイルを順次一括処理
# ==========================================
if not os.path.exists(INPUT_FOLDER):
  os.makedirs(INPUT_FOLDER)
  raise FileNotFoundError(
      f"'{INPUT_FOLDER}' フォルダが存在しないため作成しました。"
      "対象のExcelファイルを格納して再度実行してください。"
  )

if not os.path.exists(OUTPUT_FOLDER):
  os.makedirs(OUTPUT_FOLDER)

search_pattern = os.path.join(INPUT_FOLDER, "*.xlsx")
excel_files = glob.glob(search_pattern)

if not excel_files:
  raise FileNotFoundError(
      f"'{INPUT_FOLDER}' フォルダ内に対象のExcelファイル(.xlsx)が見つかりません。"
  )

print(f"対象ファイル数: {len(excel_files)}件")
print(f"ターゲットシート名: {TARGET_SHEET_NAME}\n")

# すべてのファイルを1つずつループ処理
for input_file in excel_files:
  base_name = os.path.basename(input_file)
  output_file = os.path.join(OUTPUT_FOLDER, base_name)
  print(f"処理中: {base_name} ...")

  try:
    # ------------------------------------------
    # A. データ加工処理 (Pandas)
    # ------------------------------------------
    with pd.ExcelFile(input_file) as xls:
      if TARGET_SHEET_NAME not in xls.sheet_names:
        print(
            f"  -> スキップ: シート '{TARGET_SHEET_NAME}' が存在しません。"
        )
        continue
      df = pd.read_excel(xls, sheet_name=TARGET_SHEET_NAME, header=HEADER_ROW)

    new_rows = []

    for _, row in df.iterrows():
      tool_cell = row.iloc[TOOL_COL_IDX]
      tool_dict, tool_common = parse_cell_to_blocks(tool_cell)

      # ツール一覧に属する有効な丸数字を抽出（順番通り）
      active_tool_nums = []
      for num in CIRCLED_NUMS:
        if num in tool_dict and tool_dict[num]:
          active_tool_nums.append(num)

      # ツール一覧に丸数字がない場合のフォールバック
      if not active_tool_nums:
        new_row = row.copy()
        new_rows.append(fill_empty_cells(new_row))
        continue

      # 各テキスト列のセルを事前に解析
      text_dicts = {}
      text_commons = {}
      for t_idx in TEXT_COL_INDICES:
        if t_idx < len(row):
          t_dict, t_comm = parse_cell_to_blocks(row.iloc[t_idx])
          text_dicts[t_idx] = t_dict
          text_commons[t_idx] = t_comm
        else:
          text_dicts[t_idx] = {}
          text_commons[t_idx] = []

      # ツール一覧の各丸数字をベースに行を展開
      for num in active_tool_nums:
        new_row = row.copy()
        # ツール一覧の値をセット（例: ①ツールA）
        new_row.iloc[TOOL_COL_IDX] = format_cell_value(
            num, tool_dict[num], tool_common, include_num=False
        )

        # 各テキスト列の処理
        for t_idx in TEXT_COL_INDICES:
          if t_idx < len(row):
            t_dict = text_dicts[t_idx]
            t_comm = text_commons[t_idx]

            if num in t_dict and t_dict[num]:
              new_row.iloc[t_idx] = format_cell_value(
                  num, t_dict[num], t_comm, include_num=True
              )
            elif t_comm:
              new_row.iloc[t_idx] = "\n".join(t_comm)
            else:
              new_row.iloc[t_idx] = "ー"

        # その他の空白セルを「ー」で埋める
        new_rows.append(fill_empty_cells(new_row))

    df_result = pd.DataFrame(new_rows)

    # ------------------------------------------
    # B. フォーマット維持の書き出し
    # ------------------------------------------
    with closing(openpyxl.load_workbook(input_file)) as wb:
      ws = wb[TARGET_SHEET_NAME]

      start_row = HEADER_ROW + 2
      ws.protection.sheet = False
      max_row = ws.max_row
      for row in ws.iter_rows(min_row=start_row, max_row=max_row):
        for cell in row:
          cell.value = None
          cell.protection = copy.copy(Protection(locked=False))

      for row_offset, (_, row_data) in enumerate(df_result.iterrows()):
        current_row = start_row + row_offset
        for col_idx, val in enumerate(row_data):
          cell = ws.cell(row=current_row, column=col_idx + 1)
          cell.value = val
          cell.protection = copy.copy(Protection(locked=False))

          prev_cell = ws.cell(
              row=max(start_row, current_row - 1), column=col_idx + 1
          )
          copy_cell_format(prev_cell, cell)
          cell.protection = copy.copy(Protection(locked=False))

      wb.save(output_file)
      print(f"  -> 成功: 保存先 [{output_file}]")

  except Exception as e:
    print(f"  -> エラー発生 ({base_name}): {e}")

print("\n=== すべてのファイルの処理が完了しました！ ===")
