from collections import Counter
import json
import math
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parent
INPUT_PATH = ROOT / "正解Input.xlsx"
PROCESSOR_PATH = ROOT / "excel_tool_data_procceser.py"
REPORT_PATH = ROOT / "テスト結果_100件.xlsx"


def get_category(case_id):
    if case_id <= 9:
        return "既存ベースケース"
    if case_id == 10:
        return "別字体と範囲"
    if case_id == 11:
        return "別字体と飛び番号"
    if case_id == 12:
        return "二桁番号"
    if case_id == 13:
        return "対象外番号の除外"
    if case_id == 14:
        return "テキストから番号補完"
    if case_id == 15:
        return "20番台番号"
    if case_id <= 25:
        return "別字体マッピング"
    if case_id <= 35:
        return "範囲と番号リスト"
    if case_id <= 45:
        return "対象番号と対象外番号"
    if case_id <= 55:
        return "共通テキストと混在"
    if case_id <= 65:
        return "ツール欄なし・番号補完"
    if case_id <= 85:
        return "番号11〜30"
    if case_id <= 90:
        return "不正・未登録・空欄"
    return "丸数字なし・変更なし想定"


def clean(value):
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def create_report(results):
    counts = Counter(row[2] for row in results)
    categories = Counter(row[1] for row in results)
    report = openpyxl.Workbook()
    summary = report.active
    summary.title = "集計"
    summary.append(["項目", "件数・内容"])
    summary.append(["テストケース総数", len(results)])
    for status in ("変換あり", "変更なし", "エラー"):
        summary.append([status, counts.get(status, 0)])
    summary.append(["実行対象", "excel_tool_data_procceser.py（1ケース1ブックで隔離実行）"])
    summary.append(["判定基準", "ツール一覧・テキスト1〜5を比較し、差分がなければ「変更なし」"])
    summary.append(["注意", "元データと通常の input_folder は変更せず、一時フォルダーで実行"])
    summary.append([])
    summary.append(["カテゴリ", "件数"])
    for name, count in categories.items():
        summary.append([name, count])
    summary.freeze_panes = "A2"
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 76

    detail = report.create_sheet("ケース別結果")
    detail.append([
        "ケースNo", "カテゴリ", "判定", "差分・結果", "入力ツール一覧", "出力ツール一覧",
        "入力テキスト1", "入力テキスト2", "入力テキスト3", "入力テキスト4", "入力テキスト5",
        "出力テキスト1", "出力テキスト2", "出力テキスト3", "出力テキスト4", "出力テキスト5", "例外",
    ])
    for result in results:
        detail.append(result)
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    for cell in detail[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315B55")
    widths = {"A": 11, "B": 26, "C": 14, "D": 36, "E": 34, "F": 34, "Q": 40}
    for column, width in widths.items():
        detail.column_dimensions[column].width = width
    for column in range(7, 17):
        detail.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 30
    for row in detail.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        detail.row_dimensions[row[0].row].height = 42

    summary["A1"].font = summary["B1"].font = Font(bold=True, color="FFFFFF")
    summary["A1"].fill = summary["B1"].fill = PatternFill("solid", fgColor="315B55")
    report.save(REPORT_PATH)
    return counts


def main():
    source_wb = openpyxl.load_workbook(INPUT_PATH, read_only=True, data_only=True)
    source_ws = source_wb["生徒のデータ-2"]
    headers = [source_ws.cell(5, column).value for column in range(1, 9)]
    cases = [
        [source_ws.cell(row, column).value for column in range(1, 9)]
        for row in range(6, 106)
    ]
    source_wb.close()
    case_ids = [str(row[0]) for row in cases]
    if case_ids != [str(case_id) for case_id in range(1, 101)]:
        raise ValueError("正解Input.xlsx のケース番号が1〜100で揃っていません")

    results = []
    with tempfile.TemporaryDirectory(prefix="excel-test-100-") as sandbox_name:
        sandbox = Path(sandbox_name)
        input_dir = sandbox / "input_folder"
        input_dir.mkdir()
        config = json.loads((ROOT / "config.json").read_text(encoding="utf-8"))
        config["input_folder"] = "./input_folder"
        config["output_folder"] = "./output_folder"
        (sandbox / "config.json").write_text(
            json.dumps(config, ensure_ascii=False), encoding="utf-8"
        )

        template_wb = openpyxl.load_workbook(INPUT_PATH)
        template_ws = template_wb["生徒のデータ-2"]
        template_ws.title = "RPA取込シート"
        for row in range(6, template_ws.max_row + 1):
            for column in range(1, 10):
                template_ws.cell(row, column).value = None
        template_path = sandbox / "template.xlsx"
        template_wb.save(template_path)
        template_wb.close()

        for case_id, values in enumerate(cases, start=1):
            case_path = input_dir / f"case_{case_id:03d}.xlsx"
            shutil.copyfile(template_path, case_path)
            case_wb = openpyxl.load_workbook(case_path)
            case_ws = case_wb["RPA取込シート"]
            for column, value in enumerate(values, start=1):
                case_ws.cell(6, column).value = value
            case_wb.save(case_path)
            case_wb.close()

        run = subprocess.run(
            [sys.executable, str(PROCESSOR_PATH)],
            cwd=sandbox,
            capture_output=True,
            text=True,
            check=False,
        )
        errors = {}
        for line in run.stdout.splitlines():
            if "エラー発生 (" in line:
                file_name = line.split("エラー発生 (", 1)[1].split(")", 1)[0]
                errors[file_name] = line.split(": ", 1)[-1]
        if run.returncode:
            errors["__runner__"] = (run.stderr or run.stdout)[-2000:]

        for case_id, input_row in enumerate(cases, start=1):
            file_name = f"case_{case_id:03d}.xlsx"
            before = [clean(value) for value in input_row]
            error = errors.get(file_name, errors.get("__runner__", ""))
            try:
                output_wb = openpyxl.load_workbook(
                    input_dir / file_name, read_only=True, data_only=True
                )
                output_ws = output_wb["RPA取込シート"]
                after = [clean(output_ws.cell(6, column).value) for column in range(1, 9)]
                output_wb.close()
            except Exception as exc:
                after = before[:]
                error = error or repr(exc)

            differences = [
                headers[column - 1] or f"列{column}"
                for column in range(3, 9)
                if before[column - 1] != after[column - 1]
            ]
            if error:
                status, detail = "エラー", error
            elif differences:
                status, detail = "変換あり", "差分: " + "、".join(differences)
            else:
                status = "変更なし"
                detail = "有効な丸数字を検出できず、入力値を維持"
            results.append([
                str(case_id), get_category(case_id), status, detail,
                before[2], after[2], *before[3:8], *after[3:8], error,
            ])

    counts = create_report(results)
    unchanged = [row[0] for row in results if row[2] == "変更なし"]
    errors = [row[0] for row in results if row[2] == "エラー"]
    print(f"ケース数: {len(results)}")
    print(f"結果集計: {dict(counts)}")
    print(f"変更なしケース: {', '.join(unchanged)}")
    print(f"例外ケース: {', '.join(errors)}")
    print(f"結果ブック: {REPORT_PATH}")


if __name__ == "__main__":
    main()