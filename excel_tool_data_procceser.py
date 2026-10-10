import copy
from contextlib import closing
import glob
import json
import os
import re

import openpyxl
import pandas as pd


# ==========================================
# 1. 設定ファイル (JSON) の読み込み
# ==========================================
CONFIG_FILE = "config.json"

try:
  with open(CONFIG_FILE, "r", encoding="utf-8") as f:
    config = json.load(f)
except FileNotFoundError:
  raise FileNotFoundError(
      f"設定ファイル '{CONFIG_FILE}' が見つかりません。同じ階層に配置してください。"
  )

INPUT_FOLDER = config["input_folder"]
OUTPUT_FOLDER = config["input_folder"]
TARGET_SHEET_NAME = config["target_sheet_name"]
OUTPUT_SHEET_NAME = config.get("output_sheet_name", TARGET_SHEET_NAME)
HEADER_ROW = config["header_row"]
TOOL_COL_IDX = config["tool_col_idx"]
TEXT_COL_INDICES = config["text_col_indices"]

CIRCLED_MAPPING = config["mapping"]
CIRCLED_NUMS = config["circled_nums"]
CIRCLED_TO_IDX = {num: i for i, num in enumerate(CIRCLED_NUMS)}
CIRCLED_CHARS = "".join(
    sorted({
        char
        for char in [*CIRCLED_MAPPING.keys(), *CIRCLED_NUMS]
        if len(char) == 1 and not char.isascii()
    })
)
ALL_CIRCLED_CHARS = f"[{re.escape(CIRCLED_CHARS)}]"
PREFIX_PATTERN = re.compile(
    rf"^((?:{ALL_CIRCLED_CHARS}\s*(?:[〜~-]\s*{ALL_CIRCLED_CHARS})?"
    rf"\s*[、,]?\s*)+)(?:[:：]\s*)?(.*)$"
)


# ==========================================
# 2. 丸数字標準化・解析関数
# ==========================================
def normalize_circled_number(char):
  return CIRCLED_MAPPING.get(char, char)


def expand_prefix_to_nums(prefix):
  """丸数字の指定（例: ①〜③、①③、①、③）を丸数字リストに展開する。"""
  expanded = []
  parts = re.split(r"[、,]", str(prefix))

  for part in parts:
    part = part.strip()
    if not part:
      continue

    range_match = re.search(
        rf"({ALL_CIRCLED_CHARS})\s*[〜~-]\s*({ALL_CIRCLED_CHARS})", part
    )
    if range_match:
      start_char, end_char = range_match.groups()
      start_num = normalize_circled_number(start_char)
      end_num = normalize_circled_number(end_char)
      start_idx = CIRCLED_TO_IDX.get(start_num)
      end_idx = CIRCLED_TO_IDX.get(end_num)
      if start_idx is not None and end_idx is not None and start_idx <= end_idx:
        for idx in range(start_idx, end_idx + 1):
          num = CIRCLED_NUMS[idx]
          if num not in expanded:
            expanded.append(num)
      continue

    for num_char in re.findall(ALL_CIRCLED_CHARS, part):
      num = normalize_circled_number(num_char)
      if num in CIRCLED_TO_IDX and num not in expanded:
        expanded.append(num)

  return expanded


def parse_cell_blocks(value):
  """セルを「丸数字付きブロック」と「共通テキスト」に分ける。"""
  if pd.isna(value) or not str(value).strip() or str(value).strip() == "ー":
    return [], []

  blocks = []
  common_texts = []
  lines = [line.strip() for line in str(value).splitlines() if line.strip()]

  for line in lines:
    match = PREFIX_PATTERN.match(line)
    if not match:
      common_texts.append(line)
      continue

    nums = expand_prefix_to_nums(match.group(1))
    body = match.group(2).strip()
    if nums and body:
      blocks.append((nums, body))
    else:
      common_texts.append(line)

  return blocks, common_texts


def collect_nums_from_blocks(blocks):
  nums = []
  for block_nums, _ in blocks:
    for num in block_nums:
      if num not in nums:
        nums.append(num)
  return nums


def nums_in_config_order(nums):
  return [num for num in CIRCLED_NUMS if num in nums]


def join_nums(nums):
  return "".join(nums_in_config_order(nums))


def normalize_text_cell(value, active_nums):
  blocks, common_texts = parse_cell_blocks(value)
  output_lines = []

  for block_nums, body in blocks:
    matched_nums = [num for num in nums_in_config_order(block_nums) if num in active_nums]
    if matched_nums:
      output_lines.append(f"{join_nums(matched_nums)}：{body}")

  for text in common_texts:
    output_lines.append(f"{join_nums(active_nums)}：{text}")

  return "\n".join(output_lines) if output_lines else "ー"


def normalize_tool_cell(value, active_nums):
  blocks, common_texts = parse_cell_blocks(value)
  output_lines = []

  for block_nums, body in blocks:
    matched_nums = [num for num in nums_in_config_order(block_nums) if num in active_nums]
    for num in matched_nums:
      output_lines.append(f"{num}{body}")

  if output_lines:
    return "\n".join(output_lines)

  if common_texts:
    tool_name = "\n".join(common_texts)
    return "\n".join(f"{num}{tool_name}" for num in active_nums)

  return "ー"


def fill_row_blanks(row):
  for col_idx in range(len(row)):
    val = row.iloc[col_idx]
    if pd.isna(val) or str(val).strip() == "":
      row.iloc[col_idx] = "ー"
  return row


def preserve_no_column_as_text(df):
  if df.empty:
    return df

  first_col = df.columns[0]
  df[first_col] = df[first_col].apply(
      lambda val: str(int(val))
      if isinstance(val, float) and val.is_integer()
      else str(val)
      if not pd.isna(val)
      else val
  )
  return df


def copy_cell_format(source_cell, target_cell):
  if source_cell.has_style:
    target_cell._style = copy.copy(source_cell._style)
  target_cell.number_format = source_cell.number_format
  target_cell.alignment = copy.copy(source_cell.alignment)
  target_cell.protection = copy.copy(source_cell.protection)


def normalize_row(row):
  tool_blocks, _ = parse_cell_blocks(row.iloc[TOOL_COL_IDX])
  active_nums = collect_nums_from_blocks(tool_blocks)

  if not active_nums:
    for text_col_idx in TEXT_COL_INDICES:
      if text_col_idx < len(row):
        blocks, _ = parse_cell_blocks(row.iloc[text_col_idx])
        for num in collect_nums_from_blocks(blocks):
          if num not in active_nums:
            active_nums.append(num)

  active_nums = nums_in_config_order(active_nums)
  if not active_nums:
    return row.copy()

  new_row = row.copy()
  new_row.iloc[TOOL_COL_IDX] = normalize_tool_cell(row.iloc[TOOL_COL_IDX], active_nums)

  for text_col_idx in TEXT_COL_INDICES:
    if text_col_idx < len(new_row):
      new_row.iloc[text_col_idx] = normalize_text_cell(row.iloc[text_col_idx], active_nums)

  return fill_row_blanks(new_row)


# ==========================================
# 3. 指定フォルダ内の複数ファイルを順次一括処理
# ==========================================
if not os.path.exists(INPUT_FOLDER):
  os.makedirs(INPUT_FOLDER)
  raise FileNotFoundError(
      f"'{INPUT_FOLDER}' フォルダが存在しないため作成しました。"
      "対象のExcelファイルを格納して再度実行してください。"
  )

if not os.path.exists(OUTPUT_FOLDER):
  os.makedirs(OUTPUT_FOLDER)

search_pattern = os.path.join(INPUT_FOLDER, "*.xlsx")
excel_files = glob.glob(search_pattern)

if not excel_files:
  raise FileNotFoundError(
      f"'{INPUT_FOLDER}' フォルダ内に対象のExcelファイル(.xlsx)が見つかりません。"
  )

print(f"対象ファイル数: {len(excel_files)}件")
print(f"ターゲットシート名: {TARGET_SHEET_NAME}\n")

for input_file in excel_files:
  base_name = os.path.basename(input_file)
  output_file = os.path.join(OUTPUT_FOLDER, base_name)
  print(f"処理中: {base_name} ...")

  try:
    with pd.ExcelFile(input_file) as xls:
      if TARGET_SHEET_NAME not in xls.sheet_names:
        print(f"  -> スキップ: シート '{TARGET_SHEET_NAME}' が存在しません。")
        continue
      df = pd.read_excel(xls, sheet_name=TARGET_SHEET_NAME, header=HEADER_ROW)
      df = preserve_no_column_as_text(df)

    df_result = pd.DataFrame([normalize_row(row) for _, row in df.iterrows()])

    with closing(openpyxl.load_workbook(input_file)) as wb:
      ws = wb[TARGET_SHEET_NAME]
      ws.title = OUTPUT_SHEET_NAME
      start_row = HEADER_ROW + 2
      max_row = ws.max_row

      for row in ws.iter_rows(min_row=start_row, max_row=max_row):
        for cell in row:
          cell.value = None

      for row_offset, (_, row_data) in enumerate(df_result.iterrows()):
        current_row = start_row + row_offset
        for col_idx, val in enumerate(row_data, start=1):
          cell = ws.cell(row=current_row, column=col_idx)
          source_cell = ws.cell(row=max(start_row, current_row - 1), column=col_idx)
          copy_cell_format(source_cell, cell)
          cell.value = val

      wb.save(output_file)
      print(f"  -> 成功: 保存先 [{output_file}]")

  except Exception as e:
    print(f"  -> エラー発生 ({base_name}): {e}")

print("\n=== すべてのファイルの処理が完了しました！ ===")
